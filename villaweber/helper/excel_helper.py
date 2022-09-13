from helper.configuration import Configuration
from openpyxl import load_workbook
import pandas as pd


class ExcelHelper:

    def __init__(self) -> None:
        self.configuration = Configuration('villaweber')
        self.configuration.load()
        self.path = self.configuration.ga_excel_path
        self.wb = load_workbook(self.path)

    def read_table(self, sheet_name, table=None) -> list[pd.DataFrame]:
        sheet = self.wb[sheet_name]
        tables = []

        for entry, data_boundary in sheet.tables.items():
            data = sheet[data_boundary]
            content = [[cell.value for cell in ent]  for ent in data]
            
            header = content[0]
            rest = content[1:]

            df = pd.DataFrame(rest, columns = header)
            if table is None:
                tables.append(df)
            else:
                if table == entry:
                    tables.append(df)
  
        return tables
